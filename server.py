#!/usr/bin/env python3
"""
SERVIDOR CENTRALIZADO - SERVICIO AUTOELÉCTRICO BRIBIESCA (v2.0 PRO)
Soporte Híbrido: PostgreSQL (Nube Render / Supabase / Neon) + SQLite (Local / Disco Persistente)
Auto-migración, Auto-sincronización y Persistencia Permanente
"""

import http.server
import socketserver
import json
import sqlite3
import os
import sys
import urllib.parse
import mimetypes
import socket
import traceback

PORT = int(os.environ.get('PORT', 3000))
DATABASE_URL = os.environ.get('DATABASE_URL') or os.environ.get('POSTGRES_URL')
DATA_DIR = os.environ.get('DATA_DIR', os.path.dirname(os.path.abspath(__file__)))
LOCAL_DB_FILE = os.path.join(DATA_DIR, 'taller_bribiesca.db')

def get_local_ip():
    try:
        s = socket.socket(socket.AF_INET, socket.SOCK_DGRAM)
        s.connect(("8.8.8.8", 80))
        ip = s.getsockname()[0]
        s.close()
        return ip
    except Exception:
        return "127.0.0.1"


# ==============================================================================
# GESTOR DE BASE DE DATOS (UNIFICADO POSTGRESQL + SQLITE)
# ==============================================================================

class DatabaseManager:
    def __init__(self, db_url=None, sqlite_path=None):
        self.db_url = db_url or DATABASE_URL
        self.sqlite_path = sqlite_path or LOCAL_DB_FILE
        self.is_postgres = bool(self.db_url and ('postgres://' in self.db_url or 'postgresql://' in self.db_url))
        self.db_type = 'PostgreSQL' if self.is_postgres else 'SQLite'
        self.pg_driver = None

        if self.is_postgres:
            try:
                import psycopg2
                import psycopg2.extras
                self.pg_driver = 'psycopg2'
            except ImportError:
                try:
                    import pg8000.dbapi
                    self.pg_driver = 'pg8000'
                except ImportError:
                    print("⚠️ Drivers de PostgreSQL (psycopg2/pg8000) no encontrados. Usando SQLite.")
                    self.is_postgres = False
                    self.db_type = 'SQLite'

    def get_connection(self):
        if self.is_postgres:
            return self._get_postgres_connection()
        else:
            conn = sqlite3.connect(self.sqlite_path, timeout=30.0)
            conn.row_factory = sqlite3.Row
            return conn

    def _get_postgres_connection(self):
        if self.pg_driver == 'psycopg2':
            import psycopg2
            url = self.db_url.replace('postgres://', 'postgresql://', 1)
            conn = psycopg2.connect(url, sslmode='require')
            return conn
        elif self.pg_driver == 'pg8000':
            import pg8000.dbapi
            url = urllib.parse.urlparse(self.db_url)
            conn = pg8000.dbapi.connect(
                user=url.username,
                password=url.password,
                host=url.hostname,
                port=url.port or 5432,
                database=url.path.lstrip('/'),
                ssl_context=True
            )
            return conn
        raise RuntimeError("No hay driver PostgreSQL disponible")

    def format_sql(self, sql):
        if self.is_postgres:
            return sql.replace('?', '%s')
        return sql

    def execute(self, sql, params=None):
        conn = self.get_connection()
        cur = conn.cursor()
        formatted_sql = self.format_sql(sql)
        try:
            if params:
                cur.execute(formatted_sql, params)
            else:
                cur.execute(formatted_sql)
            conn.commit()
            return cur
        except Exception as e:
            conn.rollback()
            raise e
        finally:
            conn.close()

    def fetch_all(self, sql, params=None):
        conn = self.get_connection()
        cur = conn.cursor()
        formatted_sql = self.format_sql(sql)
        try:
            if params:
                cur.execute(formatted_sql, params)
            else:
                cur.execute(formatted_sql)
            
            if self.is_postgres:
                col_names = [desc[0] for desc in cur.description]
                rows = cur.fetchall()
                return [dict(zip(col_names, row)) for row in rows]
            else:
                rows = cur.fetchall()
                return [dict(r) for r in rows]
        finally:
            conn.close()

    def fetch_one(self, sql, params=None):
        conn = self.get_connection()
        cur = conn.cursor()
        formatted_sql = self.format_sql(sql)
        try:
            if params:
                cur.execute(formatted_sql, params)
            else:
                cur.execute(formatted_sql)
            
            if self.is_postgres:
                row = cur.fetchone()
                if row:
                    col_names = [desc[0] for desc in cur.description]
                    return dict(zip(col_names, row))
                return None
            else:
                row = cur.fetchone()
                return dict(row) if row else None
        finally:
            conn.close()

    def init_schema(self):
        """Crea todas las tablas necesarias si no existen."""
        schema_statements = [
            """CREATE TABLE IF NOT EXISTS clientes (
                id_cliente TEXT PRIMARY KEY,
                nombre_completo TEXT NOT NULL,
                telefono TEXT,
                email TEXT,
                direccion TEXT,
                rfc_facturacion TEXT,
                fecha_registro TEXT,
                notas TEXT
            )""",
            """CREATE TABLE IF NOT EXISTS vehiculos (
                id_vehiculo TEXT PRIMARY KEY,
                id_cliente TEXT,
                marca TEXT,
                modelo TEXT,
                anio INTEGER,
                placa TEXT,
                color TEXT,
                numero_serie_vin TEXT,
                kilometraje_actual REAL,
                tipo_combustible TEXT,
                foto_vehiculo TEXT
            )""",
            """CREATE TABLE IF NOT EXISTS tecnicos (
                id_tecnico TEXT PRIMARY KEY,
                nombre_tecnico TEXT NOT NULL,
                telefono TEXT,
                especialidad TEXT,
                email TEXT,
                estado_activo TEXT
            )""",
            """CREATE TABLE IF NOT EXISTS servicios_catalogo (
                id_servicio TEXT PRIMARY KEY,
                categoria TEXT,
                nombre_servicio TEXT NOT NULL,
                descripcion TEXT,
                precio_sugerido REAL,
                tiempo_estimado_min INTEGER,
                activo TEXT
            )""",
            """CREATE TABLE IF NOT EXISTS inventario_refacciones (
                id_refaccion TEXT PRIMARY KEY,
                categoria TEXT,
                nombre_pieza TEXT NOT NULL,
                marca TEXT,
                stock_actual REAL,
                stock_minimo REAL,
                precio_costo REAL,
                precio_venta REAL,
                ubicacion_almacen TEXT,
                foto_pieza TEXT
            )""",
            """CREATE TABLE IF NOT EXISTS ordenes_trabajo (
                id_orden TEXT PRIMARY KEY,
                id_vehiculo TEXT,
                id_tecnico TEXT,
                fecha_ingreso TEXT,
                fecha_entrega_estimada TEXT,
                fecha_entrega_real TEXT,
                estado TEXT,
                prioridad TEXT,
                kilometraje_ingreso REAL,
                nivel_combustible TEXT,
                voltaje_bateria_v REAL,
                voltaje_alternador_v REAL,
                estado_marcha TEXT,
                codigos_falla_obd2 TEXT,
                falla_reportada TEXT,
                diagnostico_tecnico TEXT,
                trabajo_realizado TEXT,
                total_mano_obra REAL,
                total_refacciones REAL,
                subtotal REAL,
                descuento REAL,
                total_general REAL,
                anticipo REAL,
                saldo_pendiente REAL,
                metodo_pago TEXT,
                estado_pago TEXT,
                foto_entrada_frente TEXT,
                foto_tablero_testigos TEXT,
                firma_cliente_ingreso TEXT,
                firma_cliente_entrega TEXT,
                notas_internas TEXT
            )""",
            """CREATE TABLE IF NOT EXISTS detalle_servicios (
                id_detalle_servicio TEXT PRIMARY KEY,
                id_orden TEXT,
                id_servicio TEXT,
                notas_servicio TEXT,
                costo REAL
            )""",
            """CREATE TABLE IF NOT EXISTS detalle_refacciones (
                id_detalle_refaccion TEXT PRIMARY KEY,
                id_orden TEXT,
                id_refaccion TEXT,
                cantidad REAL,
                precio_unitario REAL,
                subtotal REAL
            )""",
            """CREATE TABLE IF NOT EXISTS configuracion_taller (
                id_config TEXT PRIMARY KEY,
                nombre_taller TEXT,
                eslogan TEXT,
                telefono_principal TEXT,
                whatsapp_notificaciones TEXT,
                email_contacto TEXT,
                direccion_taller TEXT,
                horario_atencion TEXT,
                logo_url TEXT,
                terminos_garantia TEXT
            )"""
        ]

        conn = self.get_connection()
        cur = conn.cursor()
        try:
            for stmt in schema_statements:
                cur.execute(self.format_sql(stmt))
            conn.commit()
            print(f"✅ Esquema de Base de Datos ({self.db_type}) inicializado correctamente.")
        except Exception as e:
            conn.rollback()
            print(f"❌ Error al inicializar esquema en {self.db_type}: {e}")
        finally:
            conn.close()

        # Verificar si la base está vacía para sembrar datos iniciales o migrar
        self.seed_if_empty()

    def seed_if_empty(self):
        try:
            row = self.fetch_one("SELECT COUNT(*) AS count FROM clientes")
            count = row.get('count', 0) if row else 0
            if count > 0:
                return  # Ya tiene datos

            print("🌱 Base de datos vacía detectada. Sembrando datos iniciales...")

            # 1. Si existe archivo local SQLite previo, intentar importar desde ahí
            if self.is_postgres and os.path.exists(self.sqlite_path):
                try:
                    local_conn = sqlite3.connect(self.sqlite_path)
                    local_conn.row_factory = sqlite3.Row
                    local_cur = local_conn.cursor()
                    tables = ['clientes', 'vehiculos', 'tecnicos', 'servicios_catalogo',
                              'inventario_refacciones', 'ordenes_trabajo', 'detalle_servicios',
                              'detalle_refacciones', 'configuracion_taller']
                    migrated_any = False
                    for tbl in tables:
                        rows = local_cur.execute(f"SELECT * FROM {tbl}").fetchall()
                        if rows:
                            migrated_any = True
                            for r in rows:
                                d = dict(r)
                                self._insert_generic_dict(tbl, d)
                    local_conn.close()
                    if migrated_any:
                        print("✨ Migración exitosa de SQLite local hacia PostgreSQL.")
                        return
                except Exception as e:
                    print(f"Aviso: No se pudo migrar desde SQLite local: {e}")

            # 2. Cargar desde data.js si no hay SQLite previo
            data_js_path = os.path.join(os.path.dirname(os.path.abspath(__file__)), 'data.js')
            if os.path.exists(data_js_path):
                with open(data_js_path, 'r', encoding='utf-8') as f:
                    text = f.read()
                start = text.find('{')
                end = text.rfind('}') + 1
                if start != -1 and end != -1:
                    initial_data = json.loads(text[start:end])
                    self.sync_full_state(initial_data)
                    print("✅ Datos semilla cargados exitosamente desde data.js")
        except Exception as e:
            print(f"Error en seed_if_empty: {e}")

    def _insert_generic_dict(self, table_name, data_dict):
        cols = list(data_dict.keys())
        placeholders = ', '.join(['?'] * len(cols))
        col_names = ', '.join(cols)
        pk = cols[0]
        pk_val = data_dict[pk]

        existing = self.fetch_one(f"SELECT 1 FROM {table_name} WHERE {pk} = ?", (pk_val,))
        if existing:
            update_clauses = ', '.join([f"{c} = ?" for c in cols if c != pk])
            update_params = [data_dict[c] for c in cols if c != pk] + [pk_val]
            self.execute(f"UPDATE {table_name} SET {update_clauses} WHERE {pk} = ?", update_params)
        else:
            insert_params = [data_dict[c] for c in cols]
            self.execute(f"INSERT INTO {table_name} ({col_names}) VALUES ({placeholders})", insert_params)

    def sync_full_state(self, state):
        """Inserta o actualiza un estado completo en la base de datos."""
        # 1. Clientes
        for c in state.get('Clientes', []):
            self.upsert_client({
                'ID_Cliente': c.get('ID_Cliente'),
                'Nombre_Completo': c.get('Nombre_Completo'),
                'Telefono': c.get('Telefono'),
                'Email': c.get('Email', ''),
                'Direccion': c.get('Direccion', ''),
                'RFC_Facturacion': c.get('RFC_Facturacion', ''),
                'Fecha_Registro': c.get('Fecha_Registro', ''),
                'Notas': c.get('Notas', '')
            })

        # 2. Vehículos
        for v in state.get('Vehiculos', []):
            self.upsert_vehicle({
                'ID_Vehiculo': v.get('ID_Vehiculo'),
                'ID_Cliente': v.get('ID_Cliente'),
                'Marca': v.get('Marca', ''),
                'Modelo': v.get('Modelo', ''),
                'Anio': v.get('Anio', 2020),
                'Placa': v.get('Placa', ''),
                'Color': v.get('Color', ''),
                'Numero_Serie_VIN': v.get('Numero_Serie_VIN', ''),
                'Kilometraje_Actual': v.get('Kilometraje_Actual', 0),
                'Tipo_Combustible': v.get('Tipo_Combustible', 'Gasolina'),
                'Foto_Vehiculo': v.get('Foto_Vehiculo', '')
            })

        # 3. Técnicos
        for t in state.get('Tecnicos', []):
            self.upsert_technician(t)

        # 4. Servicios
        for s in state.get('Servicios_Catalogo', []):
            self.upsert_service(s)

        # 5. Refacciones
        for r in state.get('Inventario_Refacciones', []):
            self.upsert_part(r)

        # 6. Órdenes
        for o in state.get('Ordenes_Trabajo', []):
            self.upsert_order(o)

    def upsert_client(self, data):
        cid = data.get('ID_Cliente') or data.get('id_cliente')
        if not cid:
            count_row = self.fetch_one("SELECT COUNT(*) AS c FROM clientes")
            cnt = (count_row.get('c') if count_row else 0) + 1
            cid = f"CLI-{str(cnt).zfill(3)}"

        params = (
            data.get('Nombre_Completo') or data.get('nombre_completo', ''),
            data.get('Telefono') or data.get('telefono', ''),
            data.get('Email') or data.get('email', ''),
            data.get('Direccion') or data.get('direccion', ''),
            data.get('RFC_Facturacion') or data.get('rfc_facturacion', ''),
            data.get('Fecha_Registro') or data.get('fecha_registro', ''),
            data.get('Notas') or data.get('notas', ''),
            cid
        )

        existing = self.fetch_one("SELECT 1 FROM clientes WHERE id_cliente = ?", (cid,))
        if existing:
            self.execute("""UPDATE clientes SET
                nombre_completo = ?, telefono = ?, email = ?, direccion = ?,
                rfc_facturacion = ?, fecha_registro = ?, notas = ?
                WHERE id_cliente = ?""", params)
        else:
            self.execute("""INSERT INTO clientes (
                nombre_completo, telefono, email, direccion,
                rfc_facturacion, fecha_registro, notas, id_cliente
            ) VALUES (?, ?, ?, ?, ?, ?, ?, ?)""", params)
        return cid

    def upsert_vehicle(self, data):
        vid = data.get('ID_Vehiculo') or data.get('id_vehiculo')
        if not vid:
            count_row = self.fetch_one("SELECT COUNT(*) AS c FROM vehiculos")
            cnt = (count_row.get('c') if count_row else 0) + 1
            vid = f"VEH-{str(cnt).zfill(3)}"

        placa = (data.get('Placa') or data.get('placa') or '').strip().upper()
        if not placa:
            placa = f"SIN-PLACA-{vid}"

        params = (
            data.get('ID_Cliente') or data.get('id_cliente'),
            data.get('Marca') or data.get('marca', ''),
            data.get('Modelo') or data.get('modelo', ''),
            data.get('Anio') or data.get('anio', 2020),
            placa,
            data.get('Color') or data.get('color', ''),
            data.get('Numero_Serie_VIN') or data.get('numero_serie_vin', ''),
            data.get('Kilometraje_Actual') or data.get('kilometraje_actual', 0),
            data.get('Tipo_Combustible') or data.get('tipo_combustible', 'Gasolina'),
            data.get('Foto_Vehiculo') or data.get('foto_vehiculo', ''),
            vid
        )

        existing = self.fetch_one("SELECT 1 FROM vehiculos WHERE id_vehiculo = ?", (vid,))
        if existing:
            self.execute("""UPDATE vehiculos SET
                id_cliente = ?, marca = ?, modelo = ?, anio = ?, placa = ?,
                color = ?, numero_serie_vin = ?, kilometraje_actual = ?,
                tipo_combustible = ?, foto_vehiculo = ?
                WHERE id_vehiculo = ?""", params)
        else:
            self.execute("""INSERT INTO vehiculos (
                id_cliente, marca, modelo, anio, placa,
                color, numero_serie_vin, kilometraje_actual,
                tipo_combustible, foto_vehiculo, id_vehiculo
            ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)""", params)
        return vid

    def upsert_technician(self, data):
        tid = data.get('ID_Tecnico') or data.get('id_tecnico')
        params = (
            data.get('Nombre_Tecnico') or data.get('nombre_tecnico', ''),
            data.get('Telefono') or data.get('telefono', ''),
            data.get('Especialidad') or data.get('especialidad', ''),
            data.get('Email') or data.get('email', ''),
            data.get('Estado_Activo') or data.get('estado_activo', 'Activo'),
            tid
        )
        existing = self.fetch_one("SELECT 1 FROM tecnicos WHERE id_tecnico = ?", (tid,))
        if existing:
            self.execute("UPDATE tecnicos SET nombre_tecnico = ?, telefono = ?, especialidad = ?, email = ?, estado_activo = ? WHERE id_tecnico = ?", params)
        else:
            self.execute("INSERT INTO tecnicos (nombre_tecnico, telefono, especialidad, email, estado_activo, id_tecnico) VALUES (?, ?, ?, ?, ?, ?)", params)

    def upsert_service(self, data):
        sid = data.get('ID_Servicio') or data.get('id_servicio')
        if not sid:
            count_row = self.fetch_one("SELECT COUNT(*) AS c FROM servicios_catalogo")
            cnt = (count_row.get('c') if count_row else 0) + 1
            sid = f"SRV-{str(cnt).zfill(3)}"

        params = (
            data.get('Categoria') or data.get('categoria', 'General'),
            data.get('Nombre_Servicio') or data.get('nombre_servicio', ''),
            data.get('Descripcion') or data.get('descripcion', ''),
            float(data.get('Precio_Sugerido') or data.get('precio_sugerido') or 0),
            int(data.get('Tiempo_Estimado_Min') or data.get('tiempo_estimado_min') or 60),
            data.get('Activo') or data.get('activo', 'Sí'),
            sid
        )
        existing = self.fetch_one("SELECT 1 FROM servicios_catalogo WHERE id_servicio = ?", (sid,))
        if existing:
            self.execute("""UPDATE servicios_catalogo SET
                categoria = ?, nombre_servicio = ?, descripcion = ?,
                precio_sugerido = ?, tiempo_estimado_min = ?, activo = ?
                WHERE id_servicio = ?""", params)
        else:
            self.execute("""INSERT INTO servicios_catalogo (
                categoria, nombre_servicio, descripcion,
                precio_sugerido, tiempo_estimado_min, activo, id_servicio
            ) VALUES (?, ?, ?, ?, ?, ?, ?)""", params)
        return sid

    def upsert_part(self, data):
        pid = data.get('ID_Refaccion') or data.get('id_refaccion')
        if not pid:
            count_row = self.fetch_one("SELECT COUNT(*) AS c FROM inventario_refacciones")
            cnt = (count_row.get('c') if count_row else 0) + 1
            pid = f"REF-{str(cnt).zfill(3)}"

        params = (
            data.get('Categoria') or data.get('categoria', 'General'),
            data.get('Nombre_Pieza') or data.get('nombre_pieza', ''),
            data.get('Marca') or data.get('marca', ''),
            float(data.get('Stock_Actual') or data.get('stock_actual') or 0),
            float(data.get('Stock_Minimo') or data.get('stock_minimo') or 2),
            float(data.get('Precio_Costo') or data.get('precio_costo') or 0),
            float(data.get('Precio_Venta') or data.get('precio_venta') or 0),
            data.get('Ubicacion_Almacen') or data.get('ubicacion_almacen', ''),
            data.get('Foto_Pieza') or data.get('foto_pieza', ''),
            pid
        )
        existing = self.fetch_one("SELECT 1 FROM inventario_refacciones WHERE id_refaccion = ?", (pid,))
        if existing:
            self.execute("""UPDATE inventario_refacciones SET
                categoria = ?, nombre_pieza = ?, marca = ?, stock_actual = ?,
                stock_minimo = ?, precio_costo = ?, precio_venta = ?,
                ubicacion_almacen = ?, foto_pieza = ?
                WHERE id_refaccion = ?""", params)
        else:
            self.execute("""INSERT INTO inventario_refacciones (
                categoria, nombre_pieza, marca, stock_actual,
                stock_minimo, precio_costo, precio_venta,
                ubicacion_almacen, foto_pieza, id_refaccion
            ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?)""", params)
        return pid

    def upsert_order(self, data):
        oid = data.get('ID_Orden') or data.get('id_orden')
        if not oid:
            count_row = self.fetch_one("SELECT COUNT(*) AS c FROM ordenes_trabajo")
            cnt = (count_row.get('c') if count_row else 0) + 1
            oid = f"ORD-{str(cnt).zfill(3)}"

        params = (
            data.get('ID_Vehiculo') or data.get('id_vehiculo'),
            data.get('ID_Tecnico') or data.get('id_tecnico'),
            data.get('Fecha_Ingreso') or data.get('fecha_ingreso', ''),
            data.get('Fecha_Entrega_Estimada') or data.get('fecha_entrega_estimada', ''),
            data.get('Fecha_Entrega_Real') or data.get('fecha_entrega_real', ''),
            data.get('Estado') or data.get('estado', 'Recepcionado'),
            data.get('Prioridad') or data.get('prioridad', 'Normal'),
            float(data.get('Kilometraje_Ingreso') or data.get('kilometraje_ingreso') or 0),
            data.get('Nivel_Combustible') or data.get('nivel_combustible', '1/2'),
            float(data.get('Voltaje_Bateria_V') or data.get('voltaje_bateria_v') or 0),
            float(data.get('Voltaje_Alternador_V') or data.get('voltaje_alternador_v') or 0),
            data.get('Estado_Marcha') or data.get('estado_marcha', 'Normal'),
            data.get('Codigos_Falla_OBD2') or data.get('codigos_falla_obd2', ''),
            data.get('Falla_Reportada') or data.get('falla_reportada', ''),
            data.get('Diagnostico_Tecnico') or data.get('diagnostico_tecnico', ''),
            data.get('Trabajo_Realizado') or data.get('trabajo_realizado', ''),
            float(data.get('Total_Mano_Obra') or data.get('total_mano_obra') or 0),
            float(data.get('Total_Refacciones') or data.get('total_refacciones') or 0),
            float(data.get('Subtotal') or data.get('subtotal') or 0),
            float(data.get('Descuento') or data.get('descuento') or 0),
            float(data.get('Total_General') or data.get('total_general') or 0),
            float(data.get('Anticipo') or data.get('anticipo') or 0),
            float(data.get('Saldo_Pendiente') or data.get('saldo_pendiente') or 0),
            data.get('Metodo_Pago') or data.get('metodo_pago', 'Efectivo'),
            data.get('Estado_Pago') or data.get('estado_pago', 'Pendiente'),
            data.get('Foto_Entrada_Frente') or data.get('foto_entrada_frente', ''),
            data.get('Foto_Tablero_Testigos') or data.get('foto_tablero_testigos', ''),
            data.get('Firma_Cliente_Ingreso') or data.get('firma_cliente_ingreso', ''),
            data.get('Firma_Cliente_Entrega') or data.get('firma_cliente_entrega', ''),
            data.get('Notas_Internas') or data.get('notas_internas', ''),
            oid
        )

        existing = self.fetch_one("SELECT 1 FROM ordenes_trabajo WHERE id_orden = ?", (oid,))
        if existing:
            self.execute("""UPDATE ordenes_trabajo SET
                id_vehiculo = ?, id_tecnico = ?, fecha_ingreso = ?, fecha_entrega_estimada = ?,
                fecha_entrega_real = ?, estado = ?, prioridad = ?, kilometraje_ingreso = ?,
                nivel_combustible = ?, voltaje_bateria_v = ?, voltaje_alternador_v = ?,
                estado_marcha = ?, codigos_falla_obd2 = ?, falla_reportada = ?,
                diagnostico_tecnico = ?, trabajo_realizado = ?, total_mano_obra = ?,
                total_refacciones = ?, subtotal = ?, descuento = ?, total_general = ?,
                anticipo = ?, saldo_pendiente = ?, metodo_pago = ?, estado_pago = ?,
                foto_entrada_frente = ?, foto_tablero_testigos = ?, firma_cliente_ingreso = ?,
                firma_cliente_entrega = ?, notas_internas = ?
                WHERE id_orden = ?""", params)
        else:
            self.execute("""INSERT INTO ordenes_trabajo (
                id_vehiculo, id_tecnico, fecha_ingreso, fecha_entrega_estimada,
                fecha_entrega_real, estado, prioridad, kilometraje_ingreso,
                nivel_combustible, voltaje_bateria_v, voltaje_alternador_v,
                estado_marcha, codigos_falla_obd2, falla_reportada,
                diagnostico_tecnico, trabajo_realizado, total_mano_obra,
                total_refacciones, subtotal, descuento, total_general,
                anticipo, saldo_pendiente, metodo_pago, estado_pago,
                foto_entrada_frente, foto_tablero_testigos, firma_cliente_ingreso,
                firma_cliente_entrega, notas_internas, id_orden
            ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)""", params)
        return oid


# Instancia Global del Gestor
db_manager = DatabaseManager()


# ==============================================================================
# MANEJADOR HTTP API REST Y ARCHIVOS ESTÁTICOS
# ==============================================================================

class CentralDatabaseHandler(http.server.BaseHTTPRequestHandler):

    def send_cors_headers(self):
        self.send_header('Access-Control-Allow-Origin', '*')
        self.send_header('Access-Control-Allow-Methods', 'GET, POST, PUT, DELETE, OPTIONS')
        self.send_header('Access-Control-Allow-Headers', 'Content-Type, Authorization')

    def do_OPTIONS(self):
        self.send_response(200)
        self.send_cors_headers()
        self.end_headers()

    def send_json(self, data, status_code=200):
        try:
            payload = json.dumps(data, ensure_ascii=False, default=str).encode('utf-8')
            self.send_response(status_code)
            self.send_header('Content-Type', 'application/json; charset=utf-8')
            self.send_header('Content-Length', str(len(payload)))
            self.send_cors_headers()
            self.end_headers()
            self.wfile.write(payload)
        except Exception as e:
            print(f"Error sending json response: {e}")

    def read_json_body(self):
        content_length = int(self.headers.get('Content-Length', 0))
        if content_length > 0:
            raw = self.rfile.read(content_length)
            try:
                return json.loads(raw.decode('utf-8'))
            except Exception as e:
                print(f"Error parsing JSON body: {e}")
                return {}
        return {}

    # === RUTAS GET ===
    def do_GET(self):
        parsed = urllib.parse.urlparse(self.path)
        path = parsed.path

        if path == '/api/status':
            try:
                orders_count = (db_manager.fetch_one("SELECT COUNT(*) AS c FROM ordenes_trabajo") or {}).get('c', 0)
                clients_count = (db_manager.fetch_one("SELECT COUNT(*) AS c FROM clientes") or {}).get('c', 0)
                parts_count = (db_manager.fetch_one("SELECT COUNT(*) AS c FROM inventario_refacciones") or {}).get('c', 0)
                self.send_json({
                    'status': 'online',
                    'database_type': db_manager.db_type,
                    'is_persistent': db_manager.is_postgres or ('/data' in LOCAL_DB_FILE),
                    'storage_info': 'Nube Permanente (PostgreSQL)' if db_manager.is_postgres else 'SQLite Local / Contenedor',
                    'counts': {
                        'orders': orders_count,
                        'clients': clients_count,
                        'parts': parts_count
                    },
                    'local_ip': get_local_ip(),
                    'port': PORT,
                    'version': '2.0-pro'
                })
            except Exception as e:
                self.send_json({'status': 'error', 'error': str(e)}, 500)
            return

        if path == '/api/all':
            self.handle_get_all_data()
            return

        if path == '/api/sync-excel':
            self.handle_sync_to_excel()
            return

        # Servir archivos estáticos
        self.serve_static_file(path)

    # === RUTAS POST ===
    def do_POST(self):
        parsed = urllib.parse.urlparse(self.path)
        path = parsed.path
        body = self.read_json_body()

        if path == '/api/orders':
            self.handle_save_order(body)
            return

        if path == '/api/clients':
            self.handle_create_client_and_vehicle(body)
            return

        if path == '/api/parts':
            self.handle_create_part(body)
            return

        if path.startswith('/api/parts/') and path.endswith('/adjust-stock'):
            parts = path.split('/')
            part_id = parts[3]
            self.handle_adjust_stock(part_id, body)
            return

        if path == '/api/services':
            self.handle_create_service(body)
            return

        if path == '/api/sync-all':
            self.handle_sync_all(body)
            return

        self.send_json({'error': 'Endpoint no encontrado'}, 404)

    # === RUTAS PUT ===
    def do_PUT(self):
        parsed = urllib.parse.urlparse(self.path)
        path = parsed.path
        body = self.read_json_body()

        if path.startswith('/api/orders/'):
            order_id = path.split('/')[3]
            body['ID_Orden'] = order_id
            self.handle_save_order(body)
            return

        self.send_json({'error': 'Endpoint no encontrado'}, 404)

    # === RUTAS DELETE ===
    def do_DELETE(self):
        parsed = urllib.parse.urlparse(self.path)
        path = parsed.path

        if path.startswith('/api/orders/'):
            order_id = path.split('/')[3]
            try:
                db_manager.execute("DELETE FROM ordenes_trabajo WHERE id_orden = ?", (order_id,))
                self.send_json({'success': True, 'id_orden': order_id, 'message': 'Orden eliminada'})
            except Exception as e:
                self.send_json({'success': False, 'error': str(e)}, 500)
            return

        self.send_json({'error': 'Endpoint no encontrado'}, 404)

    # ==========================================================================
    # API HANDLERS IMPLEMENTATION
    # ==========================================================================

    def handle_get_all_data(self):
        try:
            clientes_raw = db_manager.fetch_all("SELECT * FROM clientes")
            vehiculos_raw = db_manager.fetch_all("SELECT * FROM vehiculos")
            tecnicos_raw = db_manager.fetch_all("SELECT * FROM tecnicos")
            servicios_raw = db_manager.fetch_all("SELECT * FROM servicios_catalogo")
            refacciones_raw = db_manager.fetch_all("SELECT * FROM inventario_refacciones")
            ordenes_raw = db_manager.fetch_all("SELECT * FROM ordenes_trabajo")
            det_servicios_raw = db_manager.fetch_all("SELECT * FROM detalle_servicios")
            det_refacciones_raw = db_manager.fetch_all("SELECT * FROM detalle_refacciones")
            config_raw = db_manager.fetch_all("SELECT * FROM configuracion_taller")

            # Mapeo a CamelCase esperado por app.js
            clientes = [{
                'ID_Cliente': r.get('id_cliente'),
                'Nombre_Completo': r.get('nombre_completo'),
                'Telefono': r.get('telefono'),
                'Email': r.get('email'),
                'Direccion': r.get('direccion'),
                'RFC_Facturacion': r.get('rfc_facturacion'),
                'Fecha_Registro': r.get('fecha_registro'),
                'Notas': r.get('notas')
            } for r in clientes_raw]

            vehiculos = [{
                'ID_Vehiculo': r.get('id_vehiculo'),
                'ID_Cliente': r.get('id_cliente'),
                'Marca': r.get('marca'),
                'Modelo': r.get('modelo'),
                'Anio': r.get('anio'),
                'Placa': r.get('placa'),
                'Color': r.get('color'),
                'Numero_Serie_VIN': r.get('numero_serie_vin'),
                'Kilometraje_Actual': r.get('kilometraje_actual'),
                'Tipo_Combustible': r.get('tipo_combustible'),
                'Foto_Vehiculo': r.get('foto_vehiculo')
            } for r in vehiculos_raw]

            tecnicos = [{
                'ID_Tecnico': r.get('id_tecnico'),
                'Nombre_Tecnico': r.get('nombre_tecnico'),
                'Telefono': r.get('telefono'),
                'Especialidad': r.get('especialidad'),
                'Email': r.get('email'),
                'Estado_Activo': r.get('estado_activo')
            } for r in tecnicos_raw]

            servicios = [{
                'ID_Servicio': r.get('id_servicio'),
                'Categoria': r.get('categoria'),
                'Nombre_Servicio': r.get('nombre_servicio'),
                'Descripcion': r.get('descripcion'),
                'Precio_Sugerido': r.get('precio_sugerido'),
                'Tiempo_Estimado_Min': r.get('tiempo_estimado_min'),
                'Activo': r.get('activo')
            } for r in servicios_raw]

            refacciones = [{
                'ID_Refaccion': r.get('id_refaccion'),
                'Categoria': r.get('categoria'),
                'Nombre_Pieza': r.get('nombre_pieza'),
                'Marca': r.get('marca'),
                'Stock_Actual': r.get('stock_actual'),
                'Stock_Minimo': r.get('stock_minimo'),
                'Precio_Costo': r.get('precio_costo'),
                'Precio_Venta': r.get('precio_venta'),
                'Ubicacion_Almacen': r.get('ubicacion_almacen'),
                'Foto_Pieza': r.get('foto_pieza')
            } for r in refacciones_raw]

            ordenes = [{
                'ID_Orden': r.get('id_orden'),
                'ID_Vehiculo': r.get('id_vehiculo'),
                'ID_Tecnico': r.get('id_tecnico'),
                'Fecha_Ingreso': r.get('fecha_ingreso'),
                'Fecha_Entrega_Estimada': r.get('fecha_entrega_estimada'),
                'Fecha_Entrega_Real': r.get('fecha_entrega_real'),
                'Estado': r.get('estado'),
                'Prioridad': r.get('prioridad'),
                'Kilometraje_Ingreso': r.get('kilometraje_ingreso'),
                'Nivel_Combustible': r.get('nivel_combustible'),
                'Voltaje_Bateria_V': r.get('voltaje_bateria_v'),
                'Voltaje_Alternador_V': r.get('voltaje_alternador_v'),
                'Estado_Marcha': r.get('estado_marcha'),
                'Codigos_Falla_OBD2': r.get('codigos_falla_obd2'),
                'Falla_Reportada': r.get('falla_reportada'),
                'Diagnostico_Tecnico': r.get('diagnostico_tecnico'),
                'Trabajo_Realizado': r.get('trabajo_realizado'),
                'Total_Mano_Obra': r.get('total_mano_obra'),
                'Total_Refacciones': r.get('total_refacciones'),
                'Subtotal': r.get('subtotal'),
                'Descuento': r.get('descuento'),
                'Total_General': r.get('total_general'),
                'Anticipo': r.get('anticipo'),
                'Saldo_Pendiente': r.get('saldo_pendiente'),
                'Metodo_Pago': r.get('metodo_pago'),
                'Estado_Pago': r.get('estado_pago'),
                'Foto_Entrada_Frente': r.get('foto_entrada_frente'),
                'Foto_Tablero_Testigos': r.get('foto_tablero_testigos'),
                'Firma_Cliente_Ingreso': r.get('firma_cliente_ingreso'),
                'Firma_Cliente_Entrega': r.get('firma_cliente_entrega'),
                'Notas_Internas': r.get('notas_internas')
            } for r in ordenes_raw]

            self.send_json({
                'Clientes': clientes,
                'Vehiculos': vehiculos,
                'Tecnicos': tecnicos,
                'Servicios_Catalogo': servicios,
                'Inventario_Refacciones': refacciones,
                'Ordenes_Trabajo': ordenes,
                'Detalle_Servicios': det_servicios_raw,
                'Detalle_Refacciones': det_refacciones_raw,
                'Configuracion_Taller': config_raw
            })
        except Exception as e:
            traceback.print_exc()
            self.send_json({'error': str(e)}, 500)

    def handle_save_order(self, data):
        try:
            order_id = db_manager.upsert_order(data)
            self.send_json({
                'success': True,
                'id_orden': order_id,
                'message': f'Orden guardada exitosamente en BD ({db_manager.db_type})'
            })
        except Exception as e:
            traceback.print_exc()
            self.send_json({'success': False, 'error': str(e)}, 500)

    def handle_create_client_and_vehicle(self, data):
        try:
            cli_id = db_manager.upsert_client(data)
            veh_data = dict(data)
            veh_data['ID_Cliente'] = cli_id
            veh_id = db_manager.upsert_vehicle(veh_data)
            self.send_json({
                'success': True,
                'id_cliente': cli_id,
                'id_vehiculo': veh_id,
                'message': f'Cliente y vehículo guardados en BD ({db_manager.db_type})'
            })
        except Exception as e:
            traceback.print_exc()
            self.send_json({'success': False, 'error': str(e)}, 500)

    def handle_create_part(self, data):
        try:
            part_id = db_manager.upsert_part(data)
            self.send_json({'success': True, 'id_refaccion': part_id})
        except Exception as e:
            traceback.print_exc()
            self.send_json({'success': False, 'error': str(e)}, 500)

    def handle_adjust_stock(self, part_id, data):
        try:
            delta = int(data.get('delta', 0))
            row = db_manager.fetch_one("SELECT stock_actual, nombre_pieza FROM inventario_refacciones WHERE id_refaccion = ?", (part_id,))
            if not row:
                self.send_json({'error': 'Pieza no encontrada'}, 404)
                return

            current_stock = float(row.get('stock_actual') or 0)
            new_stock = max(0, current_stock + delta)
            db_manager.execute("UPDATE inventario_refacciones SET stock_actual = ? WHERE id_refaccion = ?", (new_stock, part_id))

            self.send_json({
                'success': True,
                'id_refaccion': part_id,
                'nombre_pieza': row.get('nombre_pieza'),
                'stock_actual': new_stock
            })
        except Exception as e:
            traceback.print_exc()
            self.send_json({'success': False, 'error': str(e)}, 500)

    def handle_create_service(self, data):
        try:
            srv_id = db_manager.upsert_service(data)
            self.send_json({'success': True, 'id_servicio': srv_id})
        except Exception as e:
            traceback.print_exc()
            self.send_json({'success': False, 'error': str(e)}, 500)

    def handle_sync_all(self, data):
        """Sincroniza un estado completo enviado desde el cliente (auto-recuperación)."""
        try:
            db_manager.sync_full_state(data)
            self.send_json({
                'success': True,
                'message': 'Estado completo sincronizado y persistido correctamente',
                'orders_synced': len(data.get('Ordenes_Trabajo', [])),
                'clients_synced': len(data.get('Clientes', []))
            })
        except Exception as e:
            traceback.print_exc()
            self.send_json({'success': False, 'error': str(e)}, 500)

    def handle_sync_to_excel(self):
        try:
            import openpyxl
            wb = openpyxl.Workbook()
            wb.remove(wb.active)

            table_sheets = [
                ('clientes', 'Clientes'),
                ('vehiculos', 'Vehiculos'),
                ('tecnicos', 'Tecnicos'),
                ('servicios_catalogo', 'Servicios_Catalogo'),
                ('inventario_refacciones', 'Inventario_Refacciones'),
                ('ordenes_trabajo', 'Ordenes_Trabajo'),
                ('detalle_servicios', 'Detalle_Servicios'),
                ('detalle_refacciones', 'Detalle_Refacciones'),
                ('configuracion_taller', 'Configuracion_Taller')
            ]

            for tbl, sheet_title in table_sheets:
                rows = db_manager.fetch_all(f"SELECT * FROM {tbl}")
                ws = wb.create_sheet(sheet_title)
                if rows:
                    ws.append(list(rows[0].keys()))
                    for r in rows:
                        ws.append(list(r.values()))

            excel_path = os.path.join(DATA_DIR, 'Servicio_Autoelectrico_Bribiesca_PRO.xlsx')
            wb.save(excel_path)
            self.send_json({'success': True, 'message': 'Excel exportado correctamente', 'path': excel_path})
        except Exception as e:
            self.send_json({'success': False, 'error': str(e)}, 500)

    # === SERVIR ARCHIVOS ESTÁTICOS ===
    def serve_static_file(self, path):
        if path == '/' or path == '':
            path = '/index.html'

        base_dir = os.path.dirname(os.path.abspath(__file__))
        clean_path = urllib.parse.unquote(path).lstrip('/')
        full_path = os.path.join(base_dir, clean_path)

        if not os.path.commonprefix([os.path.abspath(full_path), base_dir]).startswith(base_dir):
            self.send_error(403, "Acceso prohibido")
            return

        if os.path.exists(full_path) and os.path.isfile(full_path):
            mime_type, _ = mimetypes.guess_type(full_path)
            if not mime_type:
                if full_path.endswith('.js'): mime_type = 'application/javascript'
                elif full_path.endswith('.css'): mime_type = 'text/css'
                elif full_path.endswith('.html'): mime_type = 'text/html; charset=utf-8'
                elif full_path.endswith('.json'): mime_type = 'application/json'
                else: mime_type = 'application/octet-stream'

            self.send_response(200)
            self.send_header('Content-Type', mime_type)
            self.send_header('Content-Length', str(os.path.getsize(full_path)))
            self.send_cors_headers()
            self.end_headers()

            with open(full_path, 'rb') as f:
                self.wfile.write(f.read())
        else:
            self.send_error(404, "Archivo no encontrado")


# ==============================================================================
# INICIALIZACIÓN DEL SERVIDOR
# ==============================================================================

if __name__ == '__main__':
    local_ip = get_local_ip()
    print("=" * 68)
    print("⚡ SERVIDOR CENTRALIZADO - SERVICIO AUTOELÉCTRICO BRIBIESCA (v2.0 PRO)")
    print(f"🗄️  Motor de Base de Datos: {db_manager.db_type}")
    if db_manager.is_postgres:
        print(f"🌐 Conexión Nube: PostgreSQL detectado y activo")
    else:
        print(f"📁 Archivo SQLite: {LOCAL_DB_FILE}")
    print(f"🖥️  Acceso Local: http://localhost:{PORT}")
    print(f"📱 Red Taller / Dispositivos: http://{local_ip}:{PORT}")
    print("=" * 68)

    # Inicializar tablas y datos
    db_manager.init_schema()

    socketserver.ThreadingTCPServer.allow_reuse_address = True
    with socketserver.ThreadingTCPServer(("0.0.0.0", PORT), CentralDatabaseHandler) as httpd:
        try:
            httpd.serve_forever()
        except KeyboardInterrupt:
            print("\nServidor detenido correctamente.")
            sys.exit(0)
